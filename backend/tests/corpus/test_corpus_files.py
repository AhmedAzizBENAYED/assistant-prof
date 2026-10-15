import re
import zipfile
from pathlib import Path

import docx
import pytest
from corpus_gen.build import HAND_WRITTEN_MARKER, README, build_corpus, readme
from corpus_gen.documents import CORPUS, CORRUPT_NAME
from corpus_gen.package import IGNORABLE, PARA_ID, STORY_PART, TEXT_ID, W14, W
from corpus_gen.paths import DOCS_DIR
from docx.opc.exceptions import PackageNotFoundError
from lxml import etree

DOCX_NAMES = [entry.name for entry in CORPUS]
WITH_WORD_IDS = [entry.name for entry in CORPUS if entry.word_ids]
HEX_ID = re.compile(r"^[0-9A-F]{8}$")


def story_parts(path: Path) -> dict[str, etree._Element]:
    with zipfile.ZipFile(path) as package:
        return {
            name: etree.fromstring(package.read(name))
            for name in package.namelist()
            if STORY_PART.match(name)
        }


def entries(path: Path) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(path) as package:
        return [(info.filename, package.read(info)) for info in package.infolist()]


def test_committed_corpus_lists_every_generated_file() -> None:
    on_disk = {path.name for path in DOCS_DIR.glob("*.docx")}
    assert on_disk == {*DOCX_NAMES, CORRUPT_NAME}


@pytest.mark.parametrize("name", DOCX_NAMES)
def test_every_file_opens_in_python_docx(name: str) -> None:
    document = docx.Document(str(DOCS_DIR / name))
    assert any(paragraph.text.strip() for paragraph in document.paragraphs)


def test_corrupt_file_is_not_a_zip() -> None:
    assert not zipfile.is_zipfile(DOCS_DIR / CORRUPT_NAME)
    with pytest.raises(PackageNotFoundError):
        docx.Document(str(DOCS_DIR / CORRUPT_NAME))


@pytest.mark.parametrize("name", WITH_WORD_IDS)
def test_word_style_ids_on_every_paragraph(name: str) -> None:
    seen: set[str] = set()
    parts = story_parts(DOCS_DIR / name)
    assert "word/document.xml" in parts
    for part, root in parts.items():
        assert root.nsmap.get("w14") == W14, part
        assert "w14" in (root.get(IGNORABLE) or "").split(), part
        for paragraph in root.iter(f"{{{W}}}p"):
            para_id, text_id = paragraph.get(PARA_ID), paragraph.get(TEXT_ID)
            assert para_id is not None, part
            assert text_id is not None, part
            for value in (para_id, text_id):
                assert HEX_ID.match(value), value
                assert 0 < int(value, 16) < 0x80000000
            assert para_id not in seen, f"duplicated paraId {para_id} in {part}"
            seen.add(para_id)


def test_file_12_has_no_paragraph_ids() -> None:
    for root in story_parts(DOCS_DIR / "12-no-paraid.docx").values():
        for paragraph in root.iter(f"{{{W}}}p"):
            assert paragraph.get(PARA_ID) is None
            assert paragraph.get(TEXT_ID) is None


def test_headers_hold_tables_and_rtl_paragraphs() -> None:
    header_03 = story_parts(DOCS_DIR / "03-table-header-logo.docx")["word/header1.xml"]
    assert header_03.find(f"{{{W}}}tbl") is not None
    assert header_03.find(f".//{{{W}}}drawing") is not None  # the logo
    header_04 = story_parts(DOCS_DIR / "04-tunisian-bilingual-header.docx")["word/header1.xml"]
    assert len(header_04.findall(f".//{{{W}}}bidi")) == 3
    assert len(header_04.findall(f".//{{{W}}}rtl")) == 3


def test_realism_features_are_present() -> None:
    answer_lines = story_parts(DOCS_DIR / "07-answer-lines.docx")["word/document.xml"]
    leaders = {tab.get(f"{{{W}}}leader") for tab in answer_lines.iter(f"{{{W}}}tab")}
    assert {"dot", "heavy"} <= leaders
    breaks = story_parts(DOCS_DIR / "03-table-header-logo.docx")["word/document.xml"]
    assert any(br.get(f"{{{W}}}type") == "page" for br in breaks.iter(f"{{{W}}}br"))
    manual = docx.Document(str(DOCS_DIR / "02-manual-exam.docx"))
    fonts = {run.font.name for p in manual.paragraphs for run in p.runs}
    assert {"Times New Roman", "Arial", "Comic Sans MS"} <= fonts
    texts = "\n".join(p.text for p in manual.paragraphs)
    assert "\u00a0:" in texts
    assert "\u2019" in texts
    grid = docx.Document(str(DOCS_DIR / "13-dictation.docx")).tables[0]
    assert any(cell.text == "" for row in grid.rows for cell in row.cells)


def test_readme_has_one_line_per_file() -> None:
    text = (DOCS_DIR / README).read_text(encoding="utf-8")
    for name in [*DOCX_NAMES, CORRUPT_NAME]:
        rows = [line for line in text.splitlines() if f"| `{name}` |" in line]
        assert len(rows) == 1, name
        assert rows[0].count("|") == 3, f"a '|' in the description breaks the table: {name}"
    assert HAND_WRITTEN_MARKER in text


def test_readme_keeps_the_hand_written_part() -> None:
    previous = readme(None).replace("None yet.", "- `01-styled-exam.word.docx`: Word 2021.")
    assert "01-styled-exam.word.docx" in readme(previous)
    assert "None yet." in readme("no marker in this file")


def test_two_builds_are_byte_identical(fresh_corpus: Path, tmp_path: Path) -> None:
    build_corpus(tmp_path)
    names = sorted(path.name for path in fresh_corpus.iterdir())
    assert names == sorted(path.name for path in tmp_path.iterdir())
    for name in names:
        assert (fresh_corpus / name).read_bytes() == (tmp_path / name).read_bytes(), name


def test_zip_entries_have_fixed_timestamps(fresh_corpus: Path) -> None:
    for name in DOCX_NAMES:
        with zipfile.ZipFile(fresh_corpus / name) as package:
            assert {info.date_time for info in package.infolist()} == {(2026, 9, 1, 8, 0, 0)}
            core = package.read("docProps/core.xml").decode()
        assert "2026-09-01T08:00:00Z" in core


@pytest.mark.parametrize("name", DOCX_NAMES)
def test_committed_corpus_is_up_to_date(fresh_corpus: Path, name: str) -> None:
    """Compare entry contents, not zip bytes: deflate output may vary between zlib builds."""
    assert entries(DOCS_DIR / name) == entries(fresh_corpus / name), "run `just corpus`"


def test_committed_readme_is_up_to_date(fresh_corpus: Path) -> None:
    committed = (DOCS_DIR / README).read_text(encoding="utf-8")
    assert committed == readme(committed), "run `just corpus`"
