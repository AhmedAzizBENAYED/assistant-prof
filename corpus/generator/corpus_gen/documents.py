"""Builders of the python-docx corpus files (T-005a), each with the README line describing it.

The content imitates real French teachers' documents: French typography (no-break spaces,
typographic apostrophe), manual line breaks, empty paragraphs used as spacing, several points
formats, mixed fonts, typed numbers and answer lines, and the occasional typo or broken
numbering.
"""

import io
from collections.abc import Callable
from dataclasses import dataclass
from typing import Protocol, cast

import docx
from docx.document import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Length, Pt
from docx.styles.style import ParagraphStyle
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml.etree import ElementBase

from corpus_gen.logo import logo_png
from corpus_gen.typography import fr

TEXT_WIDTH = Cm(17)
"""A4 width minus 2 cm margins."""


class _Container(Protocol):
    def add_paragraph(
        self, text: str = "", style: str | ParagraphStyle | None = None
    ) -> Paragraph: ...


@dataclass(frozen=True)
class CorpusFile:
    """One generated file: its name, README line, builder, and whether Word IDs are added."""

    name: str
    description: str
    build: Callable[[], Document]
    word_ids: bool = True


# --- helpers ---------------------------------------------------------------------------------


def new_document(*, font: str | None = None, size: float | None = None) -> Document:
    """Return an A4 document with 2 cm margins, as French teachers use."""
    document = docx.Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))
    normal = document.styles["Normal"]
    assert isinstance(normal, ParagraphStyle)  # noqa: S101 - python-docx's template
    if font is not None:
        normal.font.name = font
    if size is not None:
        normal.font.size = Pt(size)
    return document


def par(container: _Container, text: str = "", style: str | None = None) -> Paragraph:
    """Add a paragraph with French typography (tabs and newlines become w:tab and w:br)."""
    return container.add_paragraph(fr(text), style)


def run(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    size: float | None = None,
    font: str | None = None,
) -> Run:
    """Add a manually formatted run, as a teacher does with the toolbar."""
    added = paragraph.add_run(fr(text))
    added.bold, added.italic, added.underline = bold, italic, underline
    if size is not None:
        added.font.size = Pt(size)
    if font is not None:
        added.font.name = font
    return added


def manual(container: _Container, *parts: tuple[str, dict[str, object]] | str) -> Paragraph:
    """Add a ``Normal`` paragraph made of runs: plain strings or ``(text, formatting)``."""
    paragraph = container.add_paragraph()
    for part in parts:
        text, formatting = (part, {}) if isinstance(part, str) else part
        run(paragraph, text, **formatting)  # type: ignore[arg-type]
    return paragraph


def tabbed(
    container: _Container,
    text: str,
    stops: list[tuple[Length, WD_TAB_ALIGNMENT, WD_TAB_LEADER]],
) -> Paragraph:
    """Add a paragraph with tab stops, e.g. answer lines drawn by a dot leader."""
    paragraph = par(container, text)
    for position, alignment, leader in stops:
        paragraph.paragraph_format.tab_stops.add_tab_stop(position, alignment, leader)
    return paragraph


def dot_leader_line(container: _Container, label: str = "") -> Paragraph:
    """Add an answer line drawn by a right tab stop with a dot leader (``w:tab leader="dot"``)."""
    return tabbed(
        container, label + "\t", [(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)]
    )


def page_break(container: _Container) -> Paragraph:
    """Add a paragraph holding only a page break, as Ctrl+Enter does."""
    paragraph = container.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def grid(table: Table, rows: list[list[str]], *, bold_header: bool = True) -> None:
    """Fill a table, one string per cell; the first row in bold."""
    for row_index, values in enumerate(rows):
        for cell, value in zip(table.rows[row_index].cells, values, strict=True):
            paragraph = cell.paragraphs[0]
            run(paragraph, value, bold=True if row_index == 0 and bold_header else None)


def header_table_first(paragraph: Paragraph, table: Table) -> None:
    """Move ``table`` before ``paragraph``: Word ends a header with an empty paragraph."""
    paragraph._p.addprevious(table._tbl)


# Children that follow w:bidi in w:pPr, and w:szCs / w:rtl / w:lang in w:rPr (schema order).
_AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)  # fmt: skip
_AFTER_SZ_CS = (
    "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl",
    "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
)  # fmt: skip
_AFTER_BCS = (
    "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
    "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color",
    "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs", *_AFTER_SZ_CS,
)  # fmt: skip
_AFTER_RTL = ("w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
_AFTER_LANG = ("w:eastAsianLayout", "w:specVanish", "w:oMath")


def element(tag: str, attributes: dict[str, str] | None = None) -> ElementBase:
    """Create a WordprocessingML element (python-docx types the result as a bare element)."""
    return cast(ElementBase, OxmlElement(tag, attributes))


def arabic(paragraph: Paragraph, text: str, *, bold: bool = False, size: float = 12) -> None:
    """Make ``paragraph`` right-to-left and add Arabic text (``w:bidi``, ``w:rtl``, cs font)."""
    paragraph._p.get_or_add_pPr().insert_element_before(element("w:bidi"), *_AFTER_BIDI)
    added = paragraph.add_run(text)
    added.bold = bold or None
    properties = added._r.get_or_add_rPr()
    properties.get_or_add_rFonts().set(qn("w:cs"), "Traditional Arabic")
    if bold:
        properties.insert_element_before(element("w:bCs"), *_AFTER_BCS)
    half_points = str(int(size * 2))
    properties.insert_element_before(element("w:szCs", {qn("w:val"): half_points}), *_AFTER_SZ_CS)
    properties.insert_element_before(element("w:rtl"), *_AFTER_RTL)
    properties.insert_element_before(element("w:lang", {qn("w:bidi"): "ar-TN"}), *_AFTER_LANG)


# --- shared exam content (files 01, 02, 14) -----------------------------------------------------

EXAM_TITLE = "Contrôle de français : le passé composé"
EXERCISE_1 = "Mets les verbes entre parenthèses au passé composé. Attention à l'auxiliaire !"
QUESTIONS_1 = (
    "Hier, nous (aller) au cinéma avec nos cousins.",
    "Les élèves (finir) leurs devoirs avant la récréation.",
    "Tu (prendre) le bus de sept heures ?",
    "Ma sœur (naître) en 2015.",
)
EXERCISE_2 = "Accorde le participe passé quand c'est nécessaire et justifie ton choix."
QUESTIONS_2 = (
    "Les fleurs que j'ai (cueillir) sont déjà fanées.",
    "Elles sont (partir) tôt ce matin.",
    "Nous avons (manger) des crêpes à la cantine.",
)
EXERCISE_3 = (
    "Raconte en huit à dix lignes ta dernière sortie scolaire.",
    "Emploie au moins cinq verbes au passé composé, dont deux avec l'auxiliaire être.",
)


# --- 01 ------------------------------------------------------------------------------------------


def styled_exam() -> Document:
    """Build file 01."""
    document = new_document(font="Times New Roman", size=12)
    section = document.sections[0]
    section.header.paragraphs[0].text = fr("Collège Ibn Khaldoun — Année scolaire 2026-2027")
    section.footer.paragraphs[0].text = fr("Contrôle n°2 — 7e année")

    par(document, EXAM_TITLE, "Title")
    dot_leader_line(document, "Nom et prénom : ")
    tabbed(
        document,
        "Classe : 7e B\tDurée : 1 heure\tNote : ...... / 20",
        [
            (Cm(7), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES),
            (TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES),
        ],
    )
    par(document)

    par(document, "Exercice 1 : Conjugue les verbes au passé composé. (6 pts)", "Heading 1")
    par(document, EXERCISE_1)
    for question in QUESTIONS_1:
        par(document, question, "List Number")
    par(document)

    # "List Number" shares one numbering: exercise 2 continues at 5, as in many real files.
    par(document, "Exercice 2 : Accorde les participes passés. (6 pts)", "Heading 1")
    par(document, EXERCISE_2)
    for question in QUESTIONS_2:
        par(document, question, "List Number")
    par(document)

    par(document, "Exercice 3 : Rédaction (8 pts)", "Heading 1")
    par(document, " ".join(EXERCISE_3))
    par(document)
    par(document)
    closing = par(document, "Bon courage !")
    closing.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    closing.runs[0].italic = True
    return document


# --- 02 ------------------------------------------------------------------------------------------


def manual_exam() -> Document:
    """Build file 02."""
    document = new_document()
    times: dict[str, object] = {"font": "Times New Roman", "size": 12}
    arial_label: dict[str, object] = {"font": "Arial", "size": 12, "bold": True, "underline": True}
    arial_note: dict[str, object] = {"font": "Arial", "size": 11, "italic": True}

    title = manual(
        document,
        ("CONTRÔLE DE FRANÇAIS\n", {"font": "Comic Sans MS", "size": 16, "bold": True}),
        ("Le passé composé", {"font": "Comic Sans MS", "size": 14}),
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    manual(document)
    manual(document, ("Nom : ____________________          Prénom : ____________________", times))
    manual(document, ("Classe : 7e ___               Date : ___ / ___ / ______", times))
    manual(document)
    manual(document)

    manual(
        document,
        ("Exercice n°1 : Conjugue les verbes au passé composé.", arial_label),
        ("                         /6", {"font": "Arial", "size": 12, "bold": True}),
    )
    manual(
        document,
        ("Mets les verbes entre parenthèses au ", arial_note),
        ("passé composé", {**arial_note, "bold": True}),
        (". Attention à l'auxiliaire !", arial_note),
    )
    for number, question in enumerate(QUESTIONS_1, start=1):
        manual(document, (f"{number}) {question}", times))
    manual(document)

    manual(
        document,
        ("Exercice n°2 : Accorde les participes passés.", arial_label),
        ("                              /6", {"font": "Arial", "size": 12, "bold": True}),
    )
    # The typo is deliberate: a "correct the spelling" instruction targets it.
    manual(document, ("Acorde le participe passé quand c'est nécessaire.", arial_note))
    for number, question in enumerate(QUESTIONS_2, start=1):
        manual(document, (f"{number}) {question}", times))
    manual(document)

    manual(
        document,
        ("Exercice n°3 : Rédaction.", arial_label),
        ("                                                    /8", {"font": "Arial", "bold": True}),
    )
    manual(document, ("\n".join(EXERCISE_3), arial_note))
    for _ in range(4):
        manual(document, ("_" * 62, times))
    manual(document)
    manual(document, (" " * 90 + "Bon courage !", {"font": "Comic Sans MS", "size": 12}))
    return document


# --- 03 ------------------------------------------------------------------------------------------


def table_header_logo() -> Document:
    """Build file 03."""
    document = new_document(font="Arial", size=11)
    section = document.sections[0]
    header = section.header
    table = header.add_table(1, 3, TEXT_WIDTH)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    school, level, date = table.rows[0].cells
    school.paragraphs[0].add_run().add_picture(io.BytesIO(logo_png()), width=Cm(1.5))
    run(school.add_paragraph(), "Collège Ibn Khaldoun\nSfax", bold=True, size=9)
    run(level.paragraphs[0], "Classe : 6e A", bold=True)
    level.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(date.paragraphs[0], "Date : ……………………")
    header_table_first(header.paragraphs[0], table)
    section.footer.paragraphs[0].text = fr("Fiche n°3 — Les déterminants et les pronoms")

    par(document, "Fiche d'exercices : les déterminants et les pronoms", "Title")
    par(document, "Partie A : Grammaire", "Heading 2")
    par(document, "Exercice 1 : Souligne les déterminants dans les phrases suivantes. /4")
    for sentence in (
        "Le chat de ma voisine dort sur le canapé.",
        "Cette année, nos professeurs organisent un voyage.",
        "Plusieurs élèves ont oublié leur cahier.",
        "Quel beau temps aujourd'hui !",
    ):
        par(document, "\u2013 " + sentence)
    par(document)
    par(document, "Exercice 2 : Remplace les groupes nominaux soulignés par un pronom. /4")
    for sentence in (
        "Je prête mon livre à Sami.",
        "Les enfants regardent les étoiles.",
        "Ma mère prépare le couscous pour mes cousins.",
    ):
        par(document, "\u2013 " + sentence)
    page_break(document)

    par(document, "Partie B : Vocabulaire", "Heading 2")
    par(document, "Exercice 3 : Trouve un synonyme pour chaque mot. /4")
    for word in ("rapide", "content", "commencer"):
        par(document, f"{word} : ........................................")
    par(document)
    par(document, "Exercice 4 : Écris le contraire de chaque adjectif. /4")
    for word in ("grand", "facile", "ancien"):
        par(document, f"{word} : ........................................")
    return document


# --- 04 ------------------------------------------------------------------------------------------


def tunisian_bilingual_header() -> Document:
    """Build file 04."""
    document = new_document(font="Times New Roman", size=12)
    section = document.sections[0]
    header = section.header
    table = header.add_table(1, 2, TEXT_WIDTH)
    french, arabic_cell = table.rows[0].cells
    for index, line in enumerate(
        (
            "République Tunisienne",
            "Ministère de l'Éducation",
            "Commissariat régional de l'éducation de Sfax 1",
        )
    ):
        paragraph = french.paragraphs[0] if index == 0 else french.add_paragraph()
        run(paragraph, line, bold=index < 2, size=11)
    arabic(arabic_cell.paragraphs[0], "الجمهورية التونسية", bold=True)
    arabic(arabic_cell.add_paragraph(), "وزارة التربية", bold=True)
    arabic(arabic_cell.add_paragraph(), "المندوبية الجهوية للتربية بصفاقس 1", size=11)
    header_table_first(header.paragraphs[0], table)

    title = manual(document, ("Devoir de contrôle n°1", {"bold": True, "size": 16}))
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tabbed(
        document,
        "Matière : Français\tNiveau : 7e année de base\tDurée : 1 h",
        [
            (Cm(8.5), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES),
            (TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES),
        ],
    )
    par(document, "Nom et prénom : ..................................   Classe : 7e B2   N° : ....")
    par(document)

    label: dict[str, object] = {"bold": True, "underline": True}
    manual(document, ("Exercice n°1 : (5 points)", label))
    par(document, "Lis attentivement le texte puis réponds aux questions.")
    manual(
        document,
        (
            "Le marché de Sfax s'éveille dès l'aube. Les marchands installent leurs étals de "
            "fruits, d'épices et de poissons. Bientôt, les voix des vendeurs se mêlent au bruit "
            "des charrettes.",
            {"italic": True},
        ),
    )
    par(document, "1) Où se passe la scène ?")
    par(document, "2) Relève deux mots qui appartiennent au champ lexical du marché.")
    par(document)
    manual(document, ("Exercice n°2 : (5 points)", label))
    par(document, "Conjugue les verbes entre parenthèses au présent de l'indicatif.")
    par(document, "1) Les marchands (installer) leurs étals.")
    par(document, "2) Nous (choisir) des oranges bien mûres.")
    par(document)
    manual(document, ("Exercice n°3 : Production écrite (10 points)", label))
    par(document, "Décris en six lignes un lieu animé de ta ville.")
    return document


# --- 07 ------------------------------------------------------------------------------------------


def answer_lines() -> Document:
    """Build file 07."""
    document = new_document(font="Arial", size=12)
    title = manual(
        document, ("Compréhension de l'écrit : Le premier jour d'école", {"bold": True, "size": 14})
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tabbed(
        document,
        "Prénom : \tClasse : \t",
        [
            (Cm(9), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS),
            (TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS),
        ],
    )
    par(document)
    par(
        document,
        "Ce matin-là, Yasmine serrait fort la main de son père. Devant le grand portail bleu, "
        "des dizaines d'enfants riaient et se bousculaient. Elle avait peur de ne connaître "
        "personne.",
    )
    par(
        document,
        "Soudain, une petite fille aux nattes rousses lui a souri et lui a tendu un bonbon. "
        "Yasmine a souri à son tour : sa première journée commençait bien.",
    )
    par(document)
    manual(document, ("Questions", {"bold": True, "underline": True}))

    par(document, "1) Qui est le personnage principal du texte ? (2 pts)")
    for _ in range(2):
        par(document, "." * 98)
    par(document, "2) Où se trouve Yasmine au début du texte ? (2 pts)")
    for _ in range(2):
        par(document, "…" * 52)
    par(document, "3) Pourquoi Yasmine a-t-elle peur ? (2 pts)")
    for _ in range(2):
        par(document, "_" * 68)
    par(document, "4) Quel geste rassure Yasmine ? Explique ta réponse. (3 pts)")
    for _ in range(3):
        dot_leader_line(document)
    par(document, "5) Imagine la suite de l'histoire en trois phrases. (3 pts)")
    for _ in range(4):
        par(document)
    tabbed(
        document,
        "Signature des parents : \t",
        [(Cm(12), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.HEAVY)],
    )
    return document


# --- 08 ------------------------------------------------------------------------------------------

MCQ: tuple[tuple[str, tuple[str, ...], bool], ...] = (
    ("Mon frère ___ oublié son cartable.", ("a", "à", "as"), False),
    ("Le chat ___ le chien dorment ensemble.", ("est", "et", "ai"), False),
    ("Ils ___ partis très tôt.", ("son", "sont", "sons"), True),
    ("Elle ___ regarde dans le miroir.", ("ce", "se", "ceux"), False),
    ("Les voisins ___ repeint leur maison.", ("on", "ont"), True),
)
"""Question, choices, and whether the choices sit on one line separated by tabs."""


def mcq() -> Document:
    """Build file 08."""
    document = new_document(font="Calibri", size=12)
    par(document, "QCM : les homophones grammaticaux", "Heading 1")
    par(document, "Entoure la bonne réponse. Chaque bonne réponse vaut 2 points.")
    par(document)
    letters = "abc"
    for number, (question, choices, one_line) in enumerate(MCQ, start=1):
        manual(document, (f"Question {number} : ", {"bold": True}), question)
        labelled = [f"{letters[i]}) {choice}" for i, choice in enumerate(choices)]
        if one_line:
            line = par(document, "\t".join(labelled))
            line.paragraph_format.left_indent = Cm(1)
        else:
            for choice in labelled:
                par(document, choice).paragraph_format.left_indent = Cm(1)
        par(document)
    manual(document, ("Total : ……… / 10", {"bold": True}))
    return document


# --- 12 ------------------------------------------------------------------------------------------


def no_paraid() -> Document:
    """Build file 12."""
    document = docx.Document()
    par(document, "Fiche de révision : les types de phrases", "Heading 1")
    par(document, "Il existe quatre types de phrases. Chacun finit par une ponctuation précise.")
    for item in (
        "La phrase déclarative donne une information. Elle se termine par un point.",
        "La phrase interrogative pose une question. Elle se termine par un point d'interrogation.",
        "La phrase exclamative exprime un sentiment fort. Elle se termine par un point "
        "d'exclamation.",
        "La phrase injonctive donne un ordre ou un conseil.",
    ):
        par(document, item, "List Bullet")
    manual(document, ("Exemples :", {"bold": True}))
    for example in (
        "Le train part à huit heures.",
        "Quand pars-tu en vacances ?",
        "Quelle belle journée !",
        "Ferme la porte, s'il te plaît.",
    ):
        par(document, example)
    manual(
        document,
        ("À retenir : ", {"bold": True}),
        "une phrase commence toujours par une majuscule.",
    )
    return document


# --- 13 ------------------------------------------------------------------------------------------


def dictation() -> Document:
    """Build file 13."""
    document = new_document(font="Times New Roman", size=13)
    par(document, "Dictée préparée : L'automne", "Heading 1")
    manual(document, ("Texte de la dictée", {"bold": True}))
    par(
        document,
        "L'automne est arrivé sans bruit. Les feuilles des platanes sont tombées une à une et "
        "ont recouvert les trottoirs d'un tapis doré.\nDans le jardin, les enfants ont ramassé "
        "des châtaignes. Le vent a soufflé toute la nuit et la pluie a frappé les volets.",
    )
    par(document)
    manual(
        document,
        ("Exercice n°1 : ", {"bold": True}),
        "Complète les phrases avec le verbe au passé composé. (5 pts)",
    )
    for sentence in (
        "Les feuilles ______________ (tomber) une à une.",
        "Les enfants ______________ (ramasser) des châtaignes.",
        "Le vent ______________ (souffler) toute la nuit.",
        "La pluie ______________ (frapper) les volets.",
        "L'automne ______________ (arriver) sans bruit.",
    ):
        par(document, sentence)
    par(document)
    manual(
        document,
        ("Exercice n°2 : ", {"bold": True}),
        "Complète le tableau de conjugaison du verbe finir. (5 pts)",
    )
    table = document.add_table(7, 3)
    table.style = "Table Grid"
    grid(
        table,
        [
            ["Pronom", "Présent", "Passé composé"],
            ["je", "finis", ""],
            ["tu", "", "as fini"],
            ["il, elle", "finit", ""],
            ["nous", "", ""],
            ["vous", "finissez", ""],
            ["ils, elles", "", "ont fini"],
        ],
    )
    par(document)
    manual(document, ("Note : ……… / 10", {"bold": True}))
    return document


# --- 14 ------------------------------------------------------------------------------------------


def answer_key() -> Document:
    """Build file 14."""
    document = new_document(font="Calibri", size=11)
    par(document, "Corrigé : contrôle de français sur le passé composé", "Title")
    par(document, "Exercice 1 (6 pts)", "Heading 1")
    par(document, "1,5 point par verbe correctement conjugué.")
    for answer in (
        "Hier, nous sommes allés au cinéma avec nos cousins.",
        "Les élèves ont fini leurs devoirs avant la récréation.",
        "Tu as pris le bus de sept heures ?",
        "Ma sœur est née en 2015.",
    ):
        par(document, answer, "List Number")
    par(document, "Exercice 2 (6 pts)", "Heading 1")
    for answer in (
        "Les fleurs que j'ai cueillies sont déjà fanées : le COD « que » est placé avant "
        "l'auxiliaire avoir.",
        "Elles sont parties tôt ce matin : accord avec le sujet (auxiliaire être).",
        "Nous avons mangé des crêpes à la cantine : pas d'accord avec l'auxiliaire avoir.",
    ):
        par(document, answer, "List Bullet")
    par(document, "Exercice 3 (8 pts)", "Heading 1")
    par(document, "Réponse libre. Voir la grille d'évaluation ci-dessous.")
    par(document, "Barème", "Heading 1")
    table = document.add_table(7, 3)
    table.style = "Table Grid"
    grid(
        table,
        [
            ["Exercice", "Critère", "Points"],
            ["Exercice 1", "Auxiliaire et participe passé corrects", "6"],
            ["Exercice 2", "Accord justifié", "6"],
            ["Exercice 3", "Respect de la consigne", "3"],
            ["Exercice 3", "Emploi du passé composé", "3"],
            ["Exercice 3", "Orthographe et ponctuation", "2"],
            ["Total", "", "20"],
        ],
    )
    for cell in table.rows[-1].cells:
        for added in cell.paragraphs[0].runs:
            added.bold = True
    return document


CORPUS: tuple[CorpusFile, ...] = (
    CorpusFile(
        "01-styled-exam.docx",
        "Clean exam: Title and Heading 1 styles, auto-numbered questions (List Number, numbering "
        "continues across exercises), dot-leader name line, tab-aligned info line, page header "
        'and footer, Times New Roman, points as "(6 pts)".',
        styled_exam,
    ),
    CorpusFile(
        "02-manual-exam.docx",
        "Same exam formatted by hand: Normal style only, manual bold/size/underline, mixed fonts "
        '(Comic Sans MS, Times New Roman, Arial), typed "1)" numbers, spaces for alignment, '
        'line breaks inside paragraphs, empty paragraphs as spacing, "Exercice n°1 :" labels, '
        'points as "/6", one deliberate typo ("Acorde").',
        manual_exam,
    ),
    CorpusFile(
        "03-table-header-logo.docx",
        "Page header made of a table (school, class, date cells) with a PNG logo; footer text; "
        'a page break between part A and part B; typed dotted lines; points as "/4".',
        table_header_logo,
    ),
    CorpusFile(
        "04-tunisian-bilingual-header.docx",
        "Tunisian bilingual page header: French (left cell) and Arabic right-to-left (right "
        "cell, w:bidi, w:rtl, complex-script font, ar-TN); manual exam body, "
        '"Exercice n°1 : (5 points)" labels, tab-aligned info line.',
        tunisian_bilingual_header,
    ),
    CorpusFile(
        "07-answer-lines.docx",
        "Reading comprehension with every kind of answer line: typed periods, typed ellipsis "
        "characters, typed underscores, empty paragraphs, and tab stops with dot and heavy "
        'leaders (w:tab leader="dot"/"heavy"); points as "(2 pts)".',
        answer_lines,
    ),
    CorpusFile(
        "08-mcq.docx",
        "Multiple-choice quiz: typed a), b), c) choices, one per paragraph or on one line "
        'separated by tabs; Heading 1 title; points as "2 points".',
        mcq,
    ),
    CorpusFile(
        "12-no-paraid.docx",
        "Plain python-docx output with no w14:paraId/w14:textId (IDs generated on import); "
        "US Letter page, default fonts, bulleted list (List Bullet).",
        no_paraid,
        word_ids=False,
    ),
    CorpusFile(
        "13-dictation.docx",
        "Dictation sheet: support text with a line break, fill-in-the-blank sentences, a "
        'conjugation grid to fill in (table with empty cells), points as "(5 pts)".',
        dictation,
    ),
    CorpusFile(
        "14-answer-key.docx",
        "Answer key: numbered and bulleted answers, Heading 1 sections, a points table with a "
        'bold total row (20 points), points as "(6 pts)" and "1,5 point".',
        answer_key,
    ),
)

CORRUPT_NAME = "corrupt.docx"
CORRUPT_DESCRIPTION = (
    "Not a valid zip: file 12 truncated to half its size (error tests). Deliberately bad."
)
CORRUPT_SOURCE = "12-no-paraid.docx"
